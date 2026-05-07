import io
import sqlite3
import pandas as pd
from itertools import combinations, product as iproduct
from datetime import datetime

from core.import_engine import DB_PATH


# ── Defaults ──────────────────────────────────────────────────────────────────

_DEFAULTS_CONFIG = [
    ("perda_min_pct",        "0.67",  "float"),
    ("perda_max_pct",        "1.70",  "float"),
    ("refilo_min_ate_3mm",   "10",    "int"  ),
    ("refilo_min_acima_3mm", "14",    "int"  ),
    ("max_comp_na_combo",    "2",     "int"  ),
    ("peso_medio_bob_pad",   "12000", "int"  ),
    ("qtd_bobinas_pad",      "1",     "int"  ),
]

_DEFAULTS_LARGURAS = [1000, 1200, 1250, 1290, 1422, 1500]


def _conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


# ── Init ──────────────────────────────────────────────────────────────────────

def init_plano_corte() -> None:
    with _conn() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS MATRIZES_PLANO_CORTE (
                codigo          TEXT,
                matriz          TEXT,
                tipo_material   TEXT,
                produto         TEXT,
                espessura       REAL,
                desenvolvimento REAL,
                importado_em    TEXT DEFAULT (datetime('now', 'localtime'))
            );
            CREATE TABLE IF NOT EXISTS CONFIG_PLANO_CORTE (
                chave  TEXT PRIMARY KEY,
                valor  TEXT NOT NULL,
                tipo   TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS LARGURAS_BOBINA (
                valor INTEGER PRIMARY KEY
            );
            CREATE TABLE IF NOT EXISTS HISTORICO_PLANO_CORTE (
                id             INTEGER PRIMARY KEY AUTOINCREMENT,
                salvo_em       TEXT DEFAULT (datetime('now', 'localtime')),
                ancora         TEXT,
                espessura      REAL,
                tipo_material  TEXT,
                largura_bobina INTEGER,
                combinacao     TEXT,
                n_ancora       INTEGER,
                total_cortes   INTEGER,
                soma_cortes_mm REAL,
                perda_mm       REAL,
                perda_pct      REAL,
                qtd_kg         REAL,
                status         TEXT,
                observacao     TEXT
            );
            CREATE TABLE IF NOT EXISTS HISTORICO_SIMULACAO (
                id               INTEGER PRIMARY KEY AUTOINCREMENT,
                salvo_em         TEXT    DEFAULT (datetime('now', 'localtime')),
                espessura        REAL    NOT NULL,
                tipo_material    TEXT    NOT NULL,
                largura_bobina   INTEGER NOT NULL,
                combinacao       TEXT    NOT NULL,
                total_cortes     INTEGER NOT NULL,
                soma_cortes_mm   REAL    NOT NULL,
                perda_mm         REAL    NOT NULL,
                perda_pct        REAL    NOT NULL,
                qtd_kg_total     REAL    NOT NULL,
                status           TEXT    NOT NULL,
                qtd_bobinas      INTEGER NOT NULL,
                peso_lote_kg     REAL    NOT NULL,
                num_fixos        INTEGER NOT NULL,
                num_sugeridos    INTEGER NOT NULL,
                observacao       TEXT    DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS HISTORICO_SIMULACAO_ITENS (
                id                 INTEGER PRIMARY KEY AUTOINCREMENT,
                simulacao_id       INTEGER NOT NULL,
                papel              TEXT    NOT NULL,
                matriz             TEXT    NOT NULL,
                codigo             TEXT,
                tipo_material      TEXT,
                descricao          TEXT,
                espessura          REAL    NOT NULL,
                desenvolvimento_mm REAL    NOT NULL,
                n_cortes           INTEGER NOT NULL,
                subtotal_mm        REAL    NOT NULL,
                qtd_kg             REAL    NOT NULL
            );
        """)
        conn.executemany(
            "INSERT OR IGNORE INTO CONFIG_PLANO_CORTE VALUES (?,?,?)",
            _DEFAULTS_CONFIG,
        )
        for v in _DEFAULTS_LARGURAS:
            conn.execute("INSERT OR IGNORE INTO LARGURAS_BOBINA VALUES (?)", (v,))


# ── Configurações ─────────────────────────────────────────────────────────────

def get_config(chave: str) -> int | float:
    with _conn() as conn:
        r = conn.execute(
            "SELECT valor, tipo FROM CONFIG_PLANO_CORTE WHERE chave = ?", (chave,)
        ).fetchone()
    return {"int": int, "float": float}[r["tipo"]](r["valor"])


def set_config(chave: str, valor) -> None:
    with _conn() as conn:
        conn.execute(
            "UPDATE CONFIG_PLANO_CORTE SET valor = ? WHERE chave = ?", (str(valor), chave)
        )


def get_all_config() -> dict:
    with _conn() as conn:
        rows = conn.execute("SELECT chave, valor, tipo FROM CONFIG_PLANO_CORTE").fetchall()
    return {r["chave"]: {"int": int, "float": float}[r["tipo"]](r["valor"]) for r in rows}


# ── Larguras de Bobina ────────────────────────────────────────────────────────

def get_larguras() -> list[int]:
    with _conn() as conn:
        rows = conn.execute("SELECT valor FROM LARGURAS_BOBINA ORDER BY valor").fetchall()
    return [r["valor"] for r in rows]


def set_larguras(lista: list[int]) -> None:
    with _conn() as conn:
        conn.execute("DELETE FROM LARGURAS_BOBINA")
        conn.executemany("INSERT INTO LARGURAS_BOBINA VALUES (?)", [(v,) for v in sorted(set(lista))])


# ── Matrizes ──────────────────────────────────────────────────────────────────

def banco_matrizes_existe() -> bool:
    with _conn() as conn:
        try:
            qtd = conn.execute("SELECT COUNT(*) FROM MATRIZES_PLANO_CORTE").fetchone()[0]
        except Exception:
            return False
    return qtd > 0


def tratar_excel_matrizes(arquivo) -> pd.DataFrame:
    df = pd.read_excel(arquivo)
    for col in ['Código', 'Matriz', 'Tipo de material', 'Produto']:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()
    df['Espessura']       = pd.to_numeric(df.get('Espessura'),       errors='coerce')
    df['Desenvolvimento'] = pd.to_numeric(df.get('Desenvolvimento'), errors='coerce')
    df = df.dropna(subset=['Espessura', 'Desenvolvimento', 'Matriz', 'Tipo de material', 'Código'])
    df = df[df['Desenvolvimento'] > 0]
    df = df[~df['Matriz'].isin(['nan', ''])]
    df = df[~df['Código'].isin(['nan', ''])]
    return df


def importar_matrizes(df: pd.DataFrame) -> int:
    df = df.rename(columns={
        'Código':           'codigo',
        'Matriz':           'matriz',
        'Tipo de material': 'tipo_material',
        'Produto':          'produto',
        'Espessura':        'espessura',
        'Desenvolvimento':  'desenvolvimento',
    })
    cols = ['codigo', 'matriz', 'tipo_material', 'produto', 'espessura', 'desenvolvimento']
    df   = df[[c for c in cols if c in df.columns]]
    with _conn() as conn:
        conn.execute("DELETE FROM MATRIZES_PLANO_CORTE")
        df.to_sql("MATRIZES_PLANO_CORTE", conn, if_exists="append", index=False)
    return len(df)


def carregar_matrizes() -> pd.DataFrame:
    with _conn() as conn:
        df = pd.read_sql("SELECT * FROM MATRIZES_PLANO_CORTE", conn)
    return df.rename(columns={
        'codigo':          'Código',
        'matriz':          'Matriz',
        'tipo_material':   'Tipo de material',
        'produto':         'Produto',
        'espessura':       'Espessura',
        'desenvolvimento': 'Desenvolvimento',
    })


def get_espessuras() -> list[float]:
    with _conn() as conn:
        rows = conn.execute(
            "SELECT DISTINCT espessura FROM MATRIZES_PLANO_CORTE ORDER BY espessura"
        ).fetchall()
    return [r["espessura"] for r in rows]


def get_tipos_material(espessura: float) -> list[str]:
    with _conn() as conn:
        rows = conn.execute(
            "SELECT DISTINCT tipo_material FROM MATRIZES_PLANO_CORTE "
            "WHERE espessura = ? ORDER BY tipo_material",
            (espessura,)
        ).fetchall()
    return [r["tipo_material"] for r in rows]


def get_matrizes_ancora(espessura: float, tipo_material: str) -> list[str]:
    with _conn() as conn:
        rows = conn.execute(
            "SELECT DISTINCT matriz FROM MATRIZES_PLANO_CORTE "
            "WHERE espessura = ? AND tipo_material = ? ORDER BY matriz",
            (espessura, tipo_material)
        ).fetchall()
    return [r["matriz"] for r in rows]


# ── Helpers internos do DataFrame ─────────────────────────────────────────────

def _dev(df: pd.DataFrame, matriz: str, espessura: float) -> float:
    mask = (df['Matriz'] == matriz) & (df['Espessura'] == espessura)
    vals = df[mask]['Desenvolvimento'].dropna()
    if vals.empty:
        raise ValueError(f"Matriz '{matriz}' não encontrada para espessura {espessura} mm.")
    return float(vals.mean())


def _cod(df: pd.DataFrame, matriz: str, espessura: float) -> str:
    mask = (df['Matriz'] == matriz) & (df['Espessura'] == espessura)
    vals = df[mask]['Código'].dropna()
    return str(vals.iloc[0]) if not vals.empty else ''


def _tipo(df: pd.DataFrame, matriz: str, espessura: float) -> str:
    mask = (df['Matriz'] == matriz) & (df['Espessura'] == espessura)
    vals = df[mask]['Tipo de material'].dropna()
    return str(vals.iloc[0]) if not vals.empty else ''


def _prod(df: pd.DataFrame, matriz: str, espessura: float) -> str:
    mask = (df['Matriz'] == matriz) & (df['Espessura'] == espessura)
    vals = df[mask]['Produto'].dropna()
    return str(vals.iloc[0]) if not vals.empty else ''


# ── Configuração carregada do banco ───────────────────────────────────────────

def _cfg() -> dict:
    try:
        return {
            'perda_min_pct':        get_config('perda_min_pct'),
            'perda_max_pct':        get_config('perda_max_pct'),
            'refilo_min_ate_3mm':   get_config('refilo_min_ate_3mm'),
            'refilo_min_acima_3mm': get_config('refilo_min_acima_3mm'),
            'max_comp_na_combo':    get_config('max_comp_na_combo'),
            'peso_medio_bob_pad':   get_config('peso_medio_bob_pad'),
            'qtd_bobinas_pad':      get_config('qtd_bobinas_pad'),
            'larguras_bobina':      get_larguras(),
        }
    except Exception:
        return {
            'perda_min_pct': 0.67, 'perda_max_pct': 1.70,
            'refilo_min_ate_3mm': 10, 'refilo_min_acima_3mm': 14,
            'max_comp_na_combo': 2, 'peso_medio_bob_pad': 12000,
            'qtd_bobinas_pad': 1, 'larguras_bobina': _DEFAULTS_LARGURAS,
        }


# ── Motor combinatório ────────────────────────────────────────────────────────

def _buscar_para_largura(
        df, dev_ancora, matriz_ancora, matrizes_comp, devs_comp,
        largura, max_comp, espessura, cfg, limite_cortes,
) -> list[dict]:
    perda_min_mm = largura * cfg['perda_min_pct'] / 100
    perda_max_mm = largura * cfg['perda_max_pct'] / 100
    refilo_min   = cfg['refilo_min_ate_3mm'] if espessura <= 3.0 else cfg['refilo_min_acima_3mm']
    max_n        = int(largura / dev_ancora)
    cod_anc      = _cod(df, matriz_ancora, espessura)
    resultados   = []

    for n_anc in range(1, max_n + 1):
        soma_anc  = dev_ancora * n_anc
        restante  = largura - soma_anc
        if restante < 0:
            break

        perda_mm     = restante
        total_cortes = n_anc
        if (perda_min_mm <= perda_mm <= perda_max_mm) and (limite_cortes is None or total_cortes <= limite_cortes):
            status = "✓ Válida" if perda_mm >= refilo_min else "Fora da regra"
            resultados.append({
                'Combinacao': f'{matriz_ancora}(x{n_anc})',
                'N_ancora': n_anc, 'Num_comp': 0, 'Total_cortes': total_cortes,
                'Detalhes': [{
                    'Matriz': matriz_ancora, 'Codigo': cod_anc,
                    'Tipo_material': _tipo(df, matriz_ancora, espessura),
                    'Descri_produto': _prod(df, matriz_ancora, espessura),
                    'Desenvolvimento_mm': dev_ancora, 'N_cortes': n_anc,
                    'Subtotal_mm': round(soma_anc, 3),
                }],
                'Soma_cortes_mm': round(soma_anc, 3),
                'Perda_mm': round(perda_mm, 3),
                'Perda_pct': round(perda_mm / largura * 100, 4),
                'Largura_bobina': largura, 'Status': status,
            })

        if not devs_comp or restante < min(devs_comp):
            continue

        indices = [i for i, d in enumerate(devs_comp) if d <= restante]
        if not indices:
            continue

        for qtd in range(1, min(max_comp, len(indices)) + 1):
            for idx_sel in combinations(indices, qtd):
                devs  = [devs_comp[i] for i in idx_sel]
                nomes = [matrizes_comp[i] for i in idx_sel]
                maxes = [max(1, int(restante / d)) for d in devs]

                for qtds in iproduct(*[range(1, mx + 1) for mx in maxes]):
                    soma_comp    = sum(d * n for d, n in zip(devs, qtds))
                    soma_total   = soma_anc + soma_comp
                    total_cortes = n_anc + sum(qtds)

                    if soma_total > largura:
                        continue
                    if limite_cortes is not None and total_cortes > limite_cortes:
                        continue

                    perda_mm = largura - soma_total
                    if not (perda_min_mm <= perda_mm <= perda_max_mm):
                        continue

                    status   = "✓ Válida" if perda_mm >= refilo_min else "Fora da regra"
                    detalhes = [{
                        'Matriz': matriz_ancora, 'Codigo': cod_anc,
                        'Tipo_material': _tipo(df, matriz_ancora, espessura),
                        'Descri_produto': _prod(df, matriz_ancora, espessura),
                        'Desenvolvimento_mm': dev_ancora, 'N_cortes': n_anc,
                        'Subtotal_mm': round(soma_anc, 3),
                    }]
                    for nome, d, n in zip(nomes, devs, qtds):
                        detalhes.append({
                            'Matriz': nome, 'Codigo': _cod(df, nome, espessura),
                            'Tipo_material': _tipo(df, nome, espessura),
                            'Descri_produto': _prod(df, nome, espessura),
                            'Desenvolvimento_mm': d, 'N_cortes': n,
                            'Subtotal_mm': round(d * n, 3),
                        })
                    comp_str = ' + '.join(f'{nm}(x{n})' for nm, n in zip(nomes, qtds))
                    resultados.append({
                        'Combinacao': f'{matriz_ancora}(x{n_anc}) + {comp_str}',
                        'N_ancora': n_anc, 'Num_comp': qtd, 'Total_cortes': total_cortes,
                        'Detalhes': detalhes,
                        'Soma_cortes_mm': round(soma_total, 3),
                        'Perda_mm': round(perda_mm, 3),
                        'Perda_pct': round(perda_mm / largura * 100, 4),
                        'Largura_bobina': largura, 'Status': status,
                    })
    return resultados


def encontrar_combinacoes(
        df, espessura, tipo_material, matriz_ancora,
        limite_cortes=None, larguras=None, max_complementares=None,
) -> tuple:
    cfg      = _cfg()
    larguras = larguras or cfg['larguras_bobina']
    max_comp = max_complementares or cfg['max_comp_na_combo']

    dev_ancora = _dev(df, matriz_ancora, espessura)
    mask_comp  = (
        (df['Espessura'] == espessura) &
        (df['Tipo de material'] == tipo_material) &
        (df['Matriz'] != matriz_ancora)
    )
    cands = (
        df[mask_comp].groupby('Matriz')['Desenvolvimento'].mean()
        .reset_index().rename(columns={'Desenvolvimento': 'dev'})
        .sort_values('dev', ascending=False)
    )
    matrizes_comp = cands['Matriz'].tolist()
    devs_comp     = cands['dev'].tolist()

    for largura in larguras:
        if dev_ancora > largura:
            continue
        res = _buscar_para_largura(
            df, dev_ancora, matriz_ancora, matrizes_comp, devs_comp,
            largura, max_comp, espessura, cfg, limite_cortes,
        )
        if res:
            return (
                pd.DataFrame(res)
                .sort_values(['Perda_pct', 'N_ancora', 'Num_comp'])
                .reset_index(drop=True),
                largura,
            )

    return pd.DataFrame(), 0


def simular_com_fixos(
        df, espessura, tipo_material, itens_fixos,
        limite_cortes=None, larguras=None, max_complementares=None,
) -> pd.DataFrame:
    cfg      = _cfg()
    larguras = larguras or cfg['larguras_bobina']
    max_comp = max_complementares or cfg['max_comp_na_combo']

    detalhes_fixos     = []
    soma_fixa          = 0.0
    total_cortes_fixos = 0

    for item in itens_fixos:
        d         = _dev(df, item['matriz'], espessura)
        soma_fixa          += d * item['n_cortes']
        total_cortes_fixos += item['n_cortes']
        detalhes_fixos.append({
            'Matriz': item['matriz'], 'Codigo': _cod(df, item['matriz'], espessura),
            'Tipo_material': _tipo(df, item['matriz'], espessura),
            'Descri_produto': _prod(df, item['matriz'], espessura),
            'Desenvolvimento_mm': d, 'N_cortes': item['n_cortes'],
            'Subtotal_mm': round(d * item['n_cortes'], 3), 'Papel': 'FIXO',
        })

    fixos_str  = ' + '.join(f"{d['Matriz']}(x{d['N_cortes']})" for d in detalhes_fixos)
    mats_fixas = {item['matriz'] for item in itens_fixos}
    mask_comp  = (
        (df['Espessura'] == espessura) &
        (df['Tipo de material'] == tipo_material) &
        (~df['Matriz'].isin(mats_fixas))
    )
    cands = (
        df[mask_comp].groupby('Matriz')['Desenvolvimento'].mean()
        .reset_index().rename(columns={'Desenvolvimento': 'dev'})
        .sort_values('dev', ascending=False)
    )
    matrizes_comp = cands['Matriz'].tolist()
    devs_comp     = cands['dev'].tolist()
    resultados    = []

    for largura in larguras:
        if soma_fixa > largura:
            continue

        perda_min_mm = largura * cfg['perda_min_pct'] / 100
        perda_max_mm = largura * cfg['perda_max_pct'] / 100
        refilo_min   = cfg['refilo_min_ate_3mm'] if espessura <= 3.0 else cfg['refilo_min_acima_3mm']
        restante     = largura - soma_fixa

        if perda_min_mm <= restante <= perda_max_mm:
            if limite_cortes is None or total_cortes_fixos <= limite_cortes:
                status = '✓ Válida' if restante >= refilo_min else 'Fora da regra'
                resultados.append({
                    'Combinacao': fixos_str + '  (sem complementar)',
                    'Largura_bobina': largura, 'Num_comp': 0,
                    'Total_cortes': total_cortes_fixos,
                    'Detalhes': [dict(d) for d in detalhes_fixos],
                    'Soma_cortes_mm': round(soma_fixa, 3),
                    'Perda_mm': round(restante, 3),
                    'Perda_pct': round(restante / largura * 100, 4),
                    'Status': status,
                })

        if restante <= 0:
            continue

        indices = [i for i, d in enumerate(devs_comp) if d <= restante]
        if not indices:
            continue

        for qtd in range(1, min(max_comp, len(indices)) + 1):
            for idx_sel in combinations(indices, qtd):
                devs  = [devs_comp[i] for i in idx_sel]
                nomes = [matrizes_comp[i] for i in idx_sel]
                maxes = [max(1, int(restante / d)) for d in devs]

                for qtds in iproduct(*[range(1, mx + 1) for mx in maxes]):
                    soma_comp    = sum(d * n for d, n in zip(devs, qtds))
                    soma_total   = soma_fixa + soma_comp
                    total_cortes = total_cortes_fixos + sum(qtds)

                    if soma_total > largura:
                        continue
                    if limite_cortes is not None and total_cortes > limite_cortes:
                        continue

                    perda_mm = largura - soma_total
                    if not (perda_min_mm <= perda_mm <= perda_max_mm):
                        continue

                    status   = '✓ Válida' if perda_mm >= refilo_min else 'Fora da regra'
                    detalhes = [dict(d) for d in detalhes_fixos]
                    for nome, dv, n in zip(nomes, devs, qtds):
                        detalhes.append({
                            'Matriz': nome, 'Codigo': _cod(df, nome, espessura),
                            'Tipo_material': _tipo(df, nome, espessura),
                            'Descri_produto': _prod(df, nome, espessura),
                            'Desenvolvimento_mm': dv, 'N_cortes': n,
                            'Subtotal_mm': round(dv * n, 3), 'Papel': 'SUGERIDO',
                        })
                    comp_str = ' + '.join(f'{nm}(x{n})' for nm, n in zip(nomes, qtds))
                    resultados.append({
                        'Combinacao': f'{fixos_str}  +  {comp_str}',
                        'Largura_bobina': largura, 'Num_comp': qtd,
                        'Total_cortes': total_cortes, 'Detalhes': detalhes,
                        'Soma_cortes_mm': round(soma_total, 3),
                        'Perda_mm': round(perda_mm, 3),
                        'Perda_pct': round(perda_mm / largura * 100, 4),
                        'Status': status,
                    })

    if not resultados:
        return pd.DataFrame()

    _ord = {'✓ Válida': 0, 'Fora da regra': 1}
    df_res        = pd.DataFrame(resultados)
    df_res['_o']  = df_res['Status'].map(_ord).fillna(2)
    return (
        df_res.sort_values(['_o', 'Perda_pct', 'Num_comp'])
        .drop(columns='_o')
        .reset_index(drop=True)
    )


# ── Cálculos de peso / KG ─────────────────────────────────────────────────────

def calcular_peso_medio_bobina(peso: float, qtd: int) -> float:
    return peso / qtd


def calcular_kg_matriz(peso_medio, largura, n_cortes, desenvolvimento, qtd_bobinas) -> float:
    return (peso_medio / largura) * (n_cortes * desenvolvimento * qtd_bobinas)


def calcular_kg_combinacao(detalhes, peso_medio, largura, qtd_bobinas) -> float:
    return round(sum(
        calcular_kg_matriz(peso_medio, largura, d['N_cortes'], d['Desenvolvimento_mm'], qtd_bobinas)
        for d in detalhes
    ), 2)


def validar_resultado(df_res: pd.DataFrame, espessura: float) -> dict:
    cfg = _cfg()
    return {
        'total':      len(df_res),
        'validas':    int((df_res['Status'] == '✓ Válida').sum()),
        'fora_regra': int((df_res['Status'] != '✓ Válida').sum()),
        'refilo_min': cfg['refilo_min_ate_3mm'] if espessura <= 3.0 else cfg['refilo_min_acima_3mm'],
    }


# ── Exportação Excel ──────────────────────────────────────────────────────────

def exportar_excel(
        df_res, largura, ancora, espessura, tipo,
        qtd_bobinas, peso_total, limite_cortes=None,
) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    cfg = _cfg()
    wb          = Workbook()
    ws_c        = wb.active
    ws_c.title  = "Combinações"
    ws_d        = wb.create_sheet("Detalhes")

    AE = "1F4E79"; AC = "BDD7EE"; VD = "E2EFDA"
    CZ = "F2F2F2"; BR = "FFFFFF"; AM = "FFF2CC"
    RX = "EDE7F6"; LJ = "FFE0B2"
    bd = Border(
        left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),  bottom=Side(style='thin', color='CCCCCC'),
    )

    def cel(ws, r, c, v, bold=False, bg=BR, fg="000000", align="left", fmt=None, wrap=False):
        cell = ws.cell(r, c, v)
        cell.font      = Font(name="Arial", size=9, bold=bold, color=fg)
        cell.fill      = PatternFill("solid", start_color=bg, end_color=bg)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        cell.border    = bd
        if fmt:
            cell.number_format = fmt
        return cell

    pm   = calcular_peso_medio_bobina(peso_total, qtd_bobinas)
    rmin = cfg['refilo_min_ate_3mm'] if espessura <= 3.0 else cfg['refilo_min_acima_3mm']

    ws_c.row_dimensions[1].height = 22
    ws_c.cell(1, 1, "PLANO DE CORTE — COMBINAÇÕES").font = Font(name="Arial", size=13, bold=True, color=AE)
    ws_c.merge_cells("A1:I1")

    parametros = [
        ("Matriz Âncora",     ancora),
        ("Espessura",         f"{espessura} mm"),
        ("Tipo de Material",  tipo),
        ("Largura da Bobina", f"{largura} mm"),
        ("Limite de Cortes",  str(limite_cortes) if limite_cortes else "Sem limite"),
        ("Refilo Mínimo",     f"{rmin} mm"),
        ("Qtd. de Bobinas",   str(qtd_bobinas)),
        ("Peso Total Lote",   f"{peso_total:,.0f} kg"),
        ("Perda Mín. / Máx.", f"{cfg['perda_min_pct']}% / {cfg['perda_max_pct']}%"),
        ("Total Combinações", len(df_res)),
    ]
    for r, (ch, vl) in enumerate(parametros, start=2):
        cel(ws_c, r, 1, ch, bold=True, bg=AC)
        cel(ws_c, r, 2, vl, bg=AC)
        ws_c.merge_cells(f"B{r}:I{r}")

    lc = len(parametros) + 3
    ws_c.row_dimensions[lc].height = 28
    for col, tit in enumerate(["#", "Combinação", "N Âncora", "Total Cortes", "Soma (mm)", "Perda (mm)", "Perda (%)", "Qtd. KG", "Status"], 1):
        cel(ws_c, lc, col, tit, bold=True, bg=AE, fg=BR, align="center")

    for i, row in df_res.iterrows():
        lr  = lc + 1 + i
        zb  = VD if i % 2 == 0 else CZ
        kg  = calcular_kg_combinacao(row['Detalhes'], pm, largura, qtd_bobinas)
        cst = VD if row['Status'] == "✓ Válida" else LJ
        cel(ws_c, lr, 1, i + 1,               bg=zb,  align="center")
        cel(ws_c, lr, 2, row['Combinacao'],    bg=zb,  wrap=True)
        cel(ws_c, lr, 3, row['N_ancora'],      bg=AM,  align="center")
        cel(ws_c, lr, 4, row['Total_cortes'],  bg=zb,  align="center")
        cel(ws_c, lr, 5, row['Soma_cortes_mm'],bg=zb,  align="right", fmt='#,##0.000')
        cel(ws_c, lr, 6, row['Perda_mm'],      bg=zb,  align="right", fmt='#,##0.000')
        cel(ws_c, lr, 7, row['Perda_pct']/100, bg=zb,  align="right", fmt='0.0000%')
        cel(ws_c, lr, 8, kg,                   bg=RX,  align="right", fmt='#,##0.00')
        cel(ws_c, lr, 9, row['Status'],        bg=cst, align="center")
        ws_c.row_dimensions[lr].height = 16

    for col, w in zip("ABCDEFGHI", [5, 52, 10, 13, 18, 13, 12, 16, 10]):
        ws_c.column_dimensions[col].width = w

    ws_d.cell(1, 1, "DETALHES POR COMBINAÇÃO").font = Font(name="Arial", size=12, bold=True, color=AE)
    ws_d.merge_cells("A1:J1")
    for col, tit in enumerate(["# Combo", "Papel", "Código", "Espessura", "Matriz", "Tipo", "Produto", "Desenv. (mm)", "Nº Cortes", "Qtd. KG"], 1):
        cel(ws_d, 2, col, tit, bold=True, bg=AE, fg=BR, align="center")

    lr = 3
    for i, row in df_res.iterrows():
        for j, det in enumerate(row['Detalhes']):
            papel   = "ÂNCORA" if j == 0 else "Complementar"
            cor_pap = AM if j == 0 else BR
            kg_m    = calcular_kg_matriz(pm, largura, det['N_cortes'], det['Desenvolvimento_mm'], qtd_bobinas)
            cel(ws_d, lr, 1,  i + 1,                          align="center")
            cel(ws_d, lr, 2,  papel,    bg=cor_pap, align="center", bold=(j == 0))
            cel(ws_d, lr, 3,  det['Codigo'],             bg=cor_pap)
            cel(ws_d, lr, 4,  espessura, bg=cor_pap, align="center", fmt='0.00')
            cel(ws_d, lr, 5,  det['Matriz'],             bg=cor_pap)
            cel(ws_d, lr, 6,  det.get('Tipo_material', ''), bg=cor_pap)
            cel(ws_d, lr, 7,  det.get('Descri_produto', ''), bg=cor_pap)
            cel(ws_d, lr, 8,  det['Desenvolvimento_mm'], bg=cor_pap, align="right", fmt='#,##0.000')
            cel(ws_d, lr, 9,  det['N_cortes'],           bg=cor_pap, align="center")
            cel(ws_d, lr, 10, kg_m,                      bg=RX,      align="right", fmt='#,##0.00')
            lr += 1

    for col, w in zip("ABCDEFGHIJ", [10, 14, 16, 11, 28, 20, 40, 22, 16, 16]):
        ws_d.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Histórico — Plano de Corte ────────────────────────────────────────────────

def salvar_hist_plano(dados: dict) -> None:
    with _conn() as conn:
        conn.execute("""
            INSERT INTO HISTORICO_PLANO_CORTE
                (ancora, espessura, tipo_material, largura_bobina,
                 combinacao, n_ancora, total_cortes, soma_cortes_mm,
                 perda_mm, perda_pct, qtd_kg, status, observacao)
            VALUES
                (:ancora, :espessura, :tipo_material, :largura_bobina,
                 :combinacao, :n_ancora, :total_cortes, :soma_cortes_mm,
                 :perda_mm, :perda_pct, :qtd_kg, :status, :observacao)
        """, dados)


def get_hist_plano() -> pd.DataFrame:
    with _conn() as conn:
        return pd.read_sql("SELECT * FROM HISTORICO_PLANO_CORTE ORDER BY salvo_em DESC", conn)


def deletar_hist_plano(id_: int) -> None:
    with _conn() as conn:
        conn.execute("DELETE FROM HISTORICO_PLANO_CORTE WHERE id = ?", (id_,))


def limpar_hist_plano() -> None:
    with _conn() as conn:
        conn.execute("DELETE FROM HISTORICO_PLANO_CORTE")


# ── Histórico — Simulação ─────────────────────────────────────────────────────

def salvar_hist_simulacao(header: dict, itens: list[dict]) -> int:
    with _conn() as conn:
        cur = conn.execute("""
            INSERT INTO HISTORICO_SIMULACAO
                (espessura, tipo_material, largura_bobina, combinacao,
                 total_cortes, soma_cortes_mm, perda_mm, perda_pct,
                 qtd_kg_total, status, qtd_bobinas, peso_lote_kg,
                 num_fixos, num_sugeridos, observacao)
            VALUES
                (:espessura, :tipo_material, :largura_bobina, :combinacao,
                 :total_cortes, :soma_cortes_mm, :perda_mm, :perda_pct,
                 :qtd_kg_total, :status, :qtd_bobinas, :peso_lote_kg,
                 :num_fixos, :num_sugeridos, :observacao)
        """, header)
        sim_id   = cur.lastrowid
        itens_db = [{**item, "simulacao_id": sim_id} for item in itens]
        conn.executemany("""
            INSERT INTO HISTORICO_SIMULACAO_ITENS
                (simulacao_id, papel, matriz, codigo, tipo_material, descricao,
                 espessura, desenvolvimento_mm, n_cortes, subtotal_mm, qtd_kg)
            VALUES
                (:simulacao_id, :papel, :matriz, :codigo, :tipo_material, :descricao,
                 :espessura, :desenvolvimento_mm, :n_cortes, :subtotal_mm, :qtd_kg)
        """, itens_db)
    return sim_id


def get_hist_simulacao() -> pd.DataFrame:
    with _conn() as conn:
        return pd.read_sql("SELECT * FROM HISTORICO_SIMULACAO ORDER BY salvo_em DESC", conn)


def get_hist_simulacao_itens(sim_id: int) -> pd.DataFrame:
    with _conn() as conn:
        return pd.read_sql(
            "SELECT * FROM HISTORICO_SIMULACAO_ITENS WHERE simulacao_id = ? ORDER BY id",
            conn, params=(sim_id,),
        )


def deletar_hist_simulacao(id_: int) -> None:
    with _conn() as conn:
        conn.execute("DELETE FROM HISTORICO_SIMULACAO_ITENS WHERE simulacao_id = ?", (id_,))
        conn.execute("DELETE FROM HISTORICO_SIMULACAO WHERE id = ?", (id_,))


def limpar_hist_simulacao() -> None:
    with _conn() as conn:
        conn.execute("DELETE FROM HISTORICO_SIMULACAO_ITENS")
        conn.execute("DELETE FROM HISTORICO_SIMULACAO")
